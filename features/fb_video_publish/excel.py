"""Bounded XLSX import: read-only row preview and one atomic manual batch."""
from __future__ import annotations

import base64
import binascii
import copy
import io
import re
import zipfile
from datetime import datetime, timedelta

from .core import BJ, DomainError, canonical, digest, id_text, normalize_batch, parse_time, plan_pairs, stamp
from .templates import render_preview, validate_template

HEADERS = ("product", "page_id", "制作模板id", "是否使用默认token", "resource_id", "desc", "发布时间", "标题")
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 200
MAX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_SCANNED_CELLS = 200000
PAIR_KEYS = {"material_id", "page_id", "config_instance_id", "config_version", "caption_template",
             "row_number", "processing_mode", "credential_mode", "title", "scheduled_at_utc"}
OPERATION = re.compile(r"(?:[a-f0-9]{32}|[a-f0-9]{8}(?:-[a-f0-9]{4}){3}-[a-f0-9]{12})")


def fail(code, message, status=400):
    raise DomainError(code, message, status)


def text(value, label, *, maximum=16000):
    if value is None:
        return ""
    if not isinstance(value, str) or any(ord(c) < 32 and c not in "\r\n\t" for c in value):
        fail("invalid_cell", label + "必须为有效文本")
    value = value.strip()
    if len(value.encode("utf-8")) > maximum:
        fail("invalid_cell", label + "内容过长")
    return value


def cell_id(value, label):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value <= 0 or value != int(value) or len(str(int(value))) > 15:
            fail("excel_id_precision", label + "可能已被Excel舍入，请设为文本并重新填写完整ID")
        value = str(int(value))
    value = text(value, label, maximum=64)
    try:
        return id_text(value)
    except DomainError:
        fail("invalid_id", label + "必须为完整的正整数ID，请使用文本格式")


def scheduled(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        when = value
    elif isinstance(value, str):
        try:
            when = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            fail("invalid_schedule", "发布时间请填写 YYYY-MM-DD HH:mm:ss（北京时间）")
        if not re.search(r"[T ]\d{2}:\d{2}", value):
            fail("invalid_schedule", "发布时间必须包含日期及小时、分钟")
    else:
        fail("invalid_schedule", "发布时间必须为Excel日期时间或日期时间文本")
    return stamp(when.replace(tzinfo=BJ) if when.tzinfo is None else when)


def skipped_row(number, pid, mid, exc):
    return {"row_number": number, "page_id": str(pid or "")[:64], "material_id": str(mid or "")[:64],
            "code": getattr(exc, "code", "invalid_row"), "message": str(exc)[:500]}


def read_workbook(filename, content_base64):
    """No formulas or external links are evaluated, and no uploaded file is saved."""
    filename = text(filename, "文件名", maximum=512)
    if not filename.lower().endswith(".xlsx") or not filename or any(c in filename for c in "/\\"):
        fail("invalid_excel_file", "请选择.xlsx文件")
    if not isinstance(content_base64, str) or len(content_base64) > ((MAX_FILE_BYTES + 2) // 3) * 4:
        fail("excel_too_large", "Excel文件不能超过2 MiB", 413)
    try:
        raw = base64.b64decode(content_base64, validate=True)
    except (ValueError, binascii.Error):
        fail("invalid_excel_file", "Excel文件内容无效")
    if not raw or len(raw) > MAX_FILE_BYTES:
        fail("excel_too_large", "Excel文件为空或超过2 MiB", 413)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 256 or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                fail("invalid_excel_file", "Excel解压内容过大，请使用下载模板")
            if len({m.filename for m in members}) != len(members):
                fail("invalid_excel_file", "Excel压缩结构无效")
            for member in members:
                if member.flag_bits & 1 or member.filename.lower().endswith("vbaproject.bin"):
                    fail("invalid_excel_file", "不支持加密或包含宏的Excel")
                if member.filename.endswith(".xml"):
                    xml = archive.read(member).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        fail("invalid_excel_file", "Excel包含不支持的XML声明")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except DomainError:
        raise
    except ImportError:
        fail("excel_parser_unavailable", "Excel解析服务暂不可用", 503)
    except Exception:
        fail("invalid_excel_file", "无法读取Excel，请使用下载模板重新保存")
    rows, errors, total = [], [], 0
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        # Inspect every populated column: limiting to I would silently hide
        # formulas/extra data in J and beyond. Bound sparse styled-cell expansion.
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        scanned = len(header)
        names = tuple(str(cell.value or "").strip() for cell in header[:8])
        if names != HEADERS or any(c.value is not None for c in header[8:]):
            fail("excel_headers_invalid", "表头必须与下载模板的8列一致，且顺序相同")
        for number, cells in enumerate(iterator, 2):
            scanned += len(cells)
            if scanned > MAX_SCANNED_CELLS:
                fail("invalid_excel_file", "工作表包含过多空白格式单元格，请使用下载模板")
            values = [c.value for c in cells]
            if number > 10001:
                fail("excel_rows_exceeded", "工作表行号过大，请删除多余空行后重试")
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            total += 1
            if total > MAX_ROWS:
                fail("excel_rows_exceeded", "单次最多导入200个数据行，请拆分文件")
            values += [None] * max(0, 8 - len(values))
            pid, mid = values[1], values[4]
            try:
                if any(c.data_type == "f" for c in cells):
                    fail("excel_formula_not_allowed", "不支持公式，请粘贴为实际值后重试")
                if any(v is not None for v in values[8:]):
                    fail("excel_extra_columns", "模板只支持8列，请移除多余数据")
                if text(values[0], "product", maximum=128).casefold() != "dramawave":
                    fail("product_not_supported", "product必填且仅支持Dramawave")
                pid, mid = cell_id(pid, "page_id"), cell_id(mid, "resource_id")
                if len(mid) > 19:
                    fail("invalid_id", "resource_id长度无效")
                mode = text(values[2], "制作模板id", maximum=64) or "原片"
                credential = text(values[3], "是否使用默认token", maximum=64) or "是"
                if mode not in {"原片", "随机模板"}:
                    fail("invalid_processing_mode", "制作模板id仅支持原片或随机模板")
                if credential not in {"是", "否"}:
                    fail("invalid_credential_mode", "是否使用默认token仅支持是或否")
                title = text(values[7], "标题", maximum=1020)
                if len(title) > 255:
                    fail("invalid_title", "标题不能超过255个字符")
                rows.append({"row_number": number, "page_id": pid, "material_id": mid,
                             "processing_mode": "original" if mode == "原片" else "random_overlay",
                             "credential_mode": "page_pool" if credential == "是" else "operator",
                             "caption_template": text(values[5], "desc"), "title": title,
                             "scheduled_at_utc": scheduled(values[6])})
            except (DomainError, ValueError, OverflowError) as exc:
                errors.append(skipped_row(number, pid, mid, exc))
    except DomainError:
        raise
    except Exception:
        fail("invalid_excel_file", "Excel工作表内容无效，请重新保存后重试")
    finally:
        workbook.close()
    if not total:
        fail("excel_empty", "Excel只有表头或没有数据行")
    return filename, rows, errors, total


def normalize_request(raw):
    keys = {"schema_version", "operation_id", "material_ids", "page_ids", "pairs", "processing_mode", "excel"}
    if not isinstance(raw, dict) or set(raw) != keys or type(raw["schema_version"]) is not int or raw["schema_version"] != 3:
        fail("invalid_request", "Excel批次字段无效")
    pairs = raw["pairs"]
    if not isinstance(pairs, list) or not 1 <= len(pairs) <= MAX_ROWS:
        fail("invalid_request", "请选择1至200个有效行")
    clean, seen_rows = [], set()
    for p in pairs:
        if not isinstance(p, dict) or set(p) != PAIR_KEYS:
            fail("invalid_request", "Excel行字段无效")
        if type(p["row_number"]) is not int or not 2 <= p["row_number"] <= 10001 or p["row_number"] in seen_rows:
            fail("invalid_request", "Excel行号无效或重复")
        if p["processing_mode"] not in {"original", "random_overlay"} or p["credential_mode"] not in {"page_pool", "operator"}:
            fail("invalid_request", "加工方式或Token策略无效")
        title = text(p["title"], "标题", maximum=1020)
        if len(title) > 255:
            fail("invalid_title", "标题不能超过255个字符")
        when = p["scheduled_at_utc"]
        if when is not None:
            when = stamp(parse_time(when))
        clean.append({**p, "title": title, "scheduled_at_utc": when})
        seen_rows.add(p["row_number"])
    if [p["row_number"] for p in clean] != sorted(seen_rows):
        fail("invalid_request", "Excel行顺序无效")
    # Reuse the existing identity, caption, pair uniqueness and scope contracts.
    base = normalize_batch({"schema_version": 2, "operation_id": raw["operation_id"],
        "material_ids": raw["material_ids"], "page_ids": raw["page_ids"], "processing_mode": "original",
        "material_schedules": {m: None for m in raw["material_ids"]} if isinstance(raw["material_ids"], list) else {},
        "pairs": [{k: p[k] for k in ("material_id", "page_id", "config_instance_id", "config_version", "caption_template")} for p in clean]})
    by_pair = {(p["material_id"], p["page_id"]): p for p in base["pairs"]}
    clean = [{**p, **by_pair[(str(p["material_id"]), str(p["page_id"]))]} for p in clean]
    mode = aggregate_mode(clean)
    meta = raw["excel"]
    if (raw["processing_mode"] != mode or not isinstance(meta, dict) or set(meta) != {"filename", "total_rows", "skipped"}
            or type(meta["total_rows"]) is not int or not 1 <= meta["total_rows"] <= MAX_ROWS
            or not isinstance(meta["skipped"], list) or len(meta["skipped"]) + len(clean) != meta["total_rows"]):
        fail("invalid_request", "Excel摘要无效")
    text(meta["filename"], "文件名", maximum=512)
    for item in meta["skipped"]:
        if (not isinstance(item, dict) or set(item) != {"row_number", "page_id", "material_id", "code", "message"}
                or type(item["row_number"]) is not int or not 2 <= item["row_number"] <= 10001 or item["row_number"] in seen_rows):
            fail("invalid_request", "跳过行信息无效")
        for key, limit in (("page_id", 256), ("material_id", 256), ("code", 128), ("message", 2000)):
            text(item[key], key, maximum=limit)
        seen_rows.add(item["row_number"])
    return {**raw, "material_ids": base["material_ids"], "page_ids": base["page_ids"], "pairs": clean, "excel": copy.deepcopy(meta)}


def aggregate_mode(pairs):
    modes = {p["processing_mode"] for p in pairs}
    return next(iter(modes)) if len(modes) == 1 else "mixed"


class ExcelService:
    def __init__(self, service):
        self.service, self.store, self.sources = service, service.store, service.sources

    def _enabled(self):
        if not self.service.excel_enabled:
            fail("excel_gate_closed", "Excel批量发布暂时关闭", 409)

    def _operator_counts(self, actor, pairs):
        ids = list(dict.fromkeys(p["page_id"] for p in pairs if p["credential_mode"] == "operator"))
        return self.sources.pages.owner_credential_counts(ids, actor["sub_user_id"]) if ids else {}

    def _resolve(self, actor, request):
        resolved = self.service.resolve(actor, request)
        if any(p["processing_mode"] == "random_overlay" for p in request["pairs"]) and not self.service.prepare_enabled:
            fail("prepare_gate_closed", "视频加工服务未启用", 409)
        counts = self._operator_counts(actor, request["pairs"])
        if any(counts.get(p["page_id"], 0) < 1 for p in request["pairs"] if p["credential_mode"] == "operator"):
            fail("operator_authorization_unavailable", "操作人已无对应Page的有效个人授权，请重新核对", 409)
        return resolved

    @staticmethod
    def facts(actor, request, resolved, configs, repeats):
        return {"purpose": "excel-publish", "actor": actor, "request": request, "resolved": resolved,
                "configs": configs, "repeats": repeats}

    def preview(self, actor, payload):
        self._enabled()
        self.service._keys(payload, {"operation_id", "filename", "content_base64"})
        if not isinstance(payload["operation_id"], str) or not OPERATION.fullmatch(payload["operation_id"]):
            fail("invalid_request", "操作标识无效")
        if not re.fullmatch(r"[1-9][0-9]{0,30}", actor["sub_user_id"]):
            fail("actor_mapping_missing", "当前操作人尚未完成唯一用户映射", 403)
        filename, rows, skipped, total = read_workbook(payload["filename"], payload["content_base64"])
        configs = self.store.configs(actor)
        ids = list(dict.fromkeys(r["page_id"] for r in rows if r["page_id"] in configs))
        pages = {p["page_id"]: p for p in self.sources.list_pages(actor, page_ids=ids)} if ids else {}
        mids = list(dict.fromkeys(r["material_id"] for r in rows if r["page_id"] in pages))
        lookup = self.sources.resolve_materials(mids, need_desc=False, enforce_blacklist=False) if mids else {"items": {}, "errors": {}}
        counts = self._operator_counts(actor, [r for r in rows if r["page_id"] in pages])
        materials, accepted, seen, now = lookup["items"], [], set(), self.store.now_fn()
        with self.store.connect() as conn:
            conn.execute("BEGIN")
            for row in rows:
                pid, mid = row["page_id"], row["material_id"]
                try:
                    if pid not in configs:
                        fail("page_not_configured", "请先配置该Page，再重新提交这些行", 409)
                    if pid not in pages:
                        fail("page_not_found", "Page不在当前账号可访问范围内", 404)
                    if pages[pid].get("blocked_reason") or not pages[pid].get("eligible_token_count"):
                        fail("page_unavailable", "Page缺少有效发布授权", 409)
                    if mid not in materials:
                        error = lookup["errors"].get(mid, {"code": "material_unavailable", "message": "素材不存在或不符合发布条件"})
                        fail(error["code"], error["message"], 409)
                    if row["credential_mode"] == "operator" and counts.get(pid, 0) < 1:
                        fail("operator_authorization_unavailable", "当前操作人没有该Page的有效个人Token，请选择默认Token或补充授权", 409)
                    if row["processing_mode"] == "random_overlay" and not self.service.prepare_enabled:
                        fail("prepare_gate_closed", "随机模板加工服务未启用", 409)
                    config = configs[pid]
                    pair = {**row, "config_instance_id": config["config_instance_id"], "config_version": config["version"],
                            "caption_template": validate_template(row["caption_template"] or config["caption_template"])}
                    candidate = {"schema_version": 3, "pairs": [pair]}
                    plan_pairs(candidate, now)
                    self.store.check_plan(conn, actor, candidate, {"pages": pages, "materials": materials})
                    render_preview(pair["caption_template"], materials[mid])
                    if (pid, mid) in seen:
                        fail("duplicate_pair", "相同Page和素材已出现在前面的有效行，本行跳过", 409)
                    accepted.append(pair)
                    seen.add((pid, mid))
                except Exception as exc:
                    if not hasattr(exc, "code") or getattr(exc, "status", 500) >= 500:
                        raise
                    skipped.append(skipped_row(row["row_number"], pid, mid, exc))
            skipped.sort(key=lambda r: r["row_number"])
            summary = {"total": total, "valid": len(accepted), "skipped": len(skipped)}
            if not accepted:
                return {"ok": True, "request": None, "items": [], "skipped": skipped, "summary": summary,
                        "preview_token": "", "repeat_required": False, "expires_at_utc": ""}
            request = normalize_request({"schema_version": 3, "operation_id": payload["operation_id"],
                "pairs": accepted, "material_ids": list(dict.fromkeys(p["material_id"] for p in accepted)),
                "page_ids": list(dict.fromkeys(p["page_id"] for p in accepted)), "processing_mode": aggregate_mode(accepted),
                "excel": {"filename": filename, "total_rows": total, "skipped": skipped}})
            resolved = {"pages": {p: pages[p] for p in request["page_ids"]},
                        "materials": {m: materials[m] for m in request["material_ids"]}}
            current_configs, repeats = self.store.check_plan(conn, actor, request, resolved)
            plans = plan_pairs(request, now)
            self.store.check_quota(conn, plans)
            facts = self.facts(actor, request, resolved, current_configs, repeats)
        times = {p["row_number"]: when for _, p, _, when, _ in plans}
        items = [{"row_number": p["row_number"], "page_id": p["page_id"], "material_id": p["material_id"],
            "page_name": pages[p["page_id"]]["page_name"], "material_name": materials[p["material_id"]]["material_name"],
            "caption": render_preview(p["caption_template"], materials[p["material_id"]]),
            "processing_mode": p["processing_mode"], "credential_mode": p["credential_mode"], "title": p["title"],
            "planned_at_utc": stamp(times[p["row_number"]]), "schedule_kind": "scheduled" if p["scheduled_at_utc"] else "immediate",
            "repeat_required": p["page_id"] + ":" + p["material_id"] in repeats} for p in request["pairs"]]
        return {"ok": True, "request": request, "items": items, "skipped": skipped, "summary": summary,
                "preview_token": self.service._token(facts), "repeat_required": bool(repeats),
                "expires_at_utc": stamp(now + timedelta(minutes=15))}

    def submit(self, actor, payload):
        self.service._keys(payload, {"request", "preview_token", "allow_repeat"})
        if type(payload["allow_repeat"]) is not bool:
            fail("invalid_request", "重复发布确认无效")
        request = normalize_request(payload["request"])
        request_hash = digest({"request": request, "allow_repeat": payload["allow_repeat"]})
        existing = self.store.existing(actor, request["operation_id"], request_hash)
        if existing:
            return existing
        self._enabled()
        if not self.service.live_enabled:
            fail("manual_live_gate_closed", "手动发布开关关闭，未创建任务", 409)
        signed = self.service._verify(payload["preview_token"])
        resolved = self._resolve(actor, request)
        with self.store.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            old = conn.execute("SELECT * FROM fb_video_batch WHERE tenant_key=? AND owner_user_id=? AND operation_id=?",
                (actor["tenant_key"], actor["user_id"], request["operation_id"])).fetchone()
            if old:
                if old["request_hash"] != request_hash:
                    fail("operation_conflict", "同一操作标识不可用于不同提交内容", 409)
                return {"ok": True, "batch_id": old["batch_id"], "operation_id": request["operation_id"],
                        "idempotent": True, "batch": self.store._batch_dto(conn, old)}
            configs, repeats = self.store.check_plan(conn, actor, request, resolved)
            if digest(self.facts(actor, request, resolved, configs, repeats)) != signed:
                fail("preview_changed", "配置、素材或历史记录发生变化，请重新核对；尚未创建任务", 409)
            frozen = copy.deepcopy(request)
            if not payload["allow_repeat"]:
                frozen["pairs"] = [p for p in request["pairs"] if p["page_id"] + ":" + p["material_id"] not in repeats]
                for p in request["pairs"]:
                    if p["page_id"] + ":" + p["material_id"] in repeats:
                        frozen["excel"]["skipped"].append(skipped_row(p["row_number"], p["page_id"], p["material_id"],
                            DomainError("repeat_not_confirmed", "历史已成功发布，本次未确认再次发布，已跳过")))
            if not frozen["pairs"]:
                fail("no_rows_selected", "没有可提交行；历史成功行需确认再次发布", 409)
            frozen["excel"]["skipped"].sort(key=lambda r: r["row_number"])
            frozen["processing_mode"] = aggregate_mode(frozen["pairs"])
            frozen["material_ids"] = list(dict.fromkeys(p["material_id"] for p in frozen["pairs"]))
            frozen["page_ids"] = list(dict.fromkeys(p["page_id"] for p in frozen["pairs"]))
            plans = plan_pairs(frozen, self.store.now_fn())
            result = self.store.freeze(conn, actor, frozen, resolved, configs, repeats, plans, request_hash=request_hash)
            conn.commit()
            return result

    def handle(self, method, suffix, actor, payload):
        if method == "POST" and suffix == "/preview":
            return self.preview(actor, payload)
        if method == "POST" and suffix == "/batches":
            return self.submit(actor, payload)
        if method == "GET":
            operation = re.fullmatch(r"/operations/([a-f0-9-]{32,36})", suffix)
            batch = re.fullmatch(r"/batches/([a-f0-9]{32})", suffix)
            if operation:
                result = self.store.existing(actor, operation[1])
                if result is None:
                    fail("operation_not_found", "尚未查询到原导入操作，请保留操作标识核对", 404)
            elif batch:
                result = self.store.get_batch(actor, batch[1])
            else:
                fail("not_found", "接口不存在", 404)
            if not result["batch"].get("excel"):
                fail("batch_not_found", "Excel批次不存在", 404)
            return result
        fail("not_found", "接口不存在", 404)
